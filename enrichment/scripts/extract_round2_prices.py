"""Extract all monthly product-price source cells, retaining exact city attribution.

The observation-city identity is checked against the annual order and actual region
header. The municipality link is an optional, separately approved city-context link.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from prepare_round2_prices import city_literal, sha

REGION_ALIASES = {
    "Ненецкий автономный округ (Архангельская область)": "Ненецкий автономный округ",
    "Архангельская область (кроме Ненецкого автономного округа)": "Архангельская область",
    "Республика Адыгея (Адыгея)": "Республика Адыгея",
    "Республика Северная Осетия-Алания": "Республика Северная Осетия — Алания",
    "Республика Татарстан (Татарстан)": "Республика Татарстан",
    "Чувашская Республика - Чувашия": "Чувашская Республика",
    "Ханты-Мансийский автономный округ - Югра (Тюменская область)": (
        "Ханты-Мансийский автономный округ — Югра"
    ),
    "Ямало-Ненецкий автономный округ (Тюменская область)": "Ямало-Ненецкий автономный округ",
    (
        "Тюменская область (кроме Ханты-Мансийского автономного округа-Югры "
        "и Ямало-Ненецкого автономного округа)"
    ): "Тюменская область",
    "Кемеровская область - Кузбасс": "Кемеровская область",
}
FEDERAL_CITIES = {
    "45000000000": (
        "Город Москва столица Российской Федерации город федерального значения",
        "Москва",
    ),
    "40000000000": ("Город Санкт-Петербург город федерального значения", "Санкт-Петербург"),
    "67000000000": ("Город федерального значения Севастополь", "Севастополь"),
}
# Exact printed spellings keyed by the source territory code, for source-list
# reconciliation only. This never approves the unresolved ISZh-to-OKTMO link.
SOURCE_LIST_ALIAS = {("54401000000", "Орёл"): "Орел"}
# Preserve the source row but quarantine the municipality context. The January
# caption/code differs from February onward, and Order440 calls 28405000000 Bezhеtsk.
SOURCE_CODE_REVIEW = {(2023, 1, "28405000000", "Вышний Волочек")}


def source_value(value):
    if value is None:
        return None, "source_blank", ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(value) or value < 0:
            raise ValueError(f"Invalid numeric source price {value}")
        return float(value), "observed" if value > 0 else "observed_zero", str(value)
    if value in {"...", "…", "-", "—"}:
        return None, "source_marker", value
    raise ValueError(f"Unknown cell content {value!r}")


def identify_city(code, label, current_region, official):
    """Exact source label and region; federal-city data stay at federal-city grain."""
    if code in FEDERAL_CITIES:
        expected, key = FEDERAL_CITIES[code]
        if label != expected:
            raise ValueError("Changed federal-city source caption")
        return key, key
    key = SOURCE_LIST_ALIAS.get((code, label), city_literal(label))
    candidates = official.loc[
        (official.city_name_key == key) & (official.region_name == current_region)
    ]
    if len(candidates) != 1:
        raise ValueError(
            f"Unattributed source city: code={code}, label={label!r}, region={current_region!r}"
        )
    return key, current_region


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo-root", type=Path, required=True)
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--bridge-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError("Use a new extraction directory")
    args.output.mkdir(parents=True)
    lists_path = args.bridge_dir / "official_base_cities.csv"
    approved_path = args.bridge_dir / "bridge_approved_evidence.csv"
    candidates_path = args.bridge_dir / "bridge_candidates.csv"
    lists = pd.read_csv(lists_path, dtype={"region_code": str})
    approved = pd.read_csv(approved_path, dtype={"territory_id": str, "source_city_oktmo_11": str})
    candidates = pd.read_csv(candidates_path, dtype={"territory_id": str})
    temporal = candidates.loc[candidates.bridge_status == "temporal_review"]
    annual_path = args.repo_root / "data/derived/enrichment/municipal_crosswalk_2024-01-01.csv"
    regions = set(pd.read_csv(annual_path).region_name)
    con = duckdb.connect()
    con.execute("""CREATE TABLE cells (
        year INTEGER, month INTEGER, source_city_code_literal VARCHAR, product_code INTEGER,
        value DOUBLE, value_status VARCHAR, source_raw_value VARCHAR,
        source_row INTEGER, source_column INTEGER)""")
    product_rows, cities, qa, footnotes, source_files = [], [], [], [], []
    for year in (2023, 2024):
        path = args.raw_dir / f"sred_potreb_cen_{year}.xlsx"
        source_files.append(
            {
                "year": year,
                "filename": path.name,
                "sha256": sha(path),
                "url": "https://rosstat.gov.ru/storage/mediabank/" + path.name,
                "unit_header": "на конец периода, рублей за единицу измерения",
            }
        )
        workbook = load_workbook(path, read_only=True, data_only=True)
        expected_sheets = [f"{month:02}({year})" for month in range(1, 13)]
        if workbook.sheetnames[1:] != expected_sheets:
            raise ValueError("Unexpected monthly sheet inventory")
        official = lists.loc[lists.year == year]
        for month, sheet_name in enumerate(expected_sheets, 1):
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            top = [next(iterator) for _ in range(5)]
            if top[2][:2] != ("Код территории", "Наименование территории"):
                raise ValueError("Unexpected workbook row/column schema")
            if top[1][5] != "на конец периода, рублей за единицу измерения":
                raise ValueError("Unexpected unit/period semantics")
            codes, labels = top[2][2:], top[3][2:]
            if len(codes) != len(set(codes)) or not all(isinstance(code, int) for code in codes):
                raise ValueError("Duplicate/absent product code header")
            if not all(isinstance(label, str) and label for label in labels):
                raise ValueError("Missing product name")
            for col, (code, label) in enumerate(zip(codes, labels, strict=True), 3):
                product_rows.append(
                    {
                        "year": year,
                        "month": month,
                        "product_code": code,
                        "source_product_label": label,
                        "source_column": col,
                        "source_column_letter": get_column_letter(col),
                    }
                )
            region = None
            chunk, city_codes, noncity_rows = [], set(), 0
            for row_number, row in enumerate(iterator, 6):
                code, label = row[:2]
                if not isinstance(code, int):
                    if any(value is not None for value in row):
                        footnotes.append(
                            {
                                "year": year,
                                "month": month,
                                "source_row": row_number,
                                "source_values": [str(value) for value in row if value is not None],
                            }
                        )
                    continue
                source_code = str(code)  # Never pad a numeric source code implicitly.
                if code == 643 or "федеральный округ" in str(label):
                    noncity_rows += 1
                    region = None
                    continue
                region_candidate = REGION_ALIASES.get(label, label)
                if region_candidate in regions and source_code not in FEDERAL_CITIES:
                    region = region_candidate
                    noncity_rows += 1
                    continue
                key, city_region = identify_city(source_code, label, region, official)
                if source_code in city_codes:
                    raise ValueError("Duplicate source city code within month")
                city_codes.add(source_code)
                link = approved.loc[
                    (approved.year == year)
                    & (approved.city_name_key == key)
                    & (approved.region_name == city_region)
                ]
                if len(link) > 1:
                    raise ValueError("Ambiguous approved city bridge")
                temporal_link = temporal.loc[
                    (temporal.year == year)
                    & (temporal.city_name_key == key)
                    & (temporal.region_name == city_region)
                ]
                if len(temporal_link) > 1 or (len(temporal_link) and len(link)):
                    raise ValueError("Conflicting temporal city bridge")
                record = {
                    "year": year,
                    "month": month,
                    "source_city_code_literal": source_code,
                    "source_city_label": label,
                    "source_region_label": region,
                    "city_name_key": key,
                    "region_name": city_region,
                    "source_row": row_number,
                    "territory_id": None,
                    "source_city_oktmo_11": None,
                    "in_primary_cohort": False,
                    "bridge_status": "not_approved",
                    "source_code_review_reason": "",
                    "temporal_review_reason": "",
                    "candidate_territory_id": None,
                }
                needs_code_review = (year, month, source_code, label) in SOURCE_CODE_REVIEW
                if needs_code_review:
                    record.update(
                        {
                            "bridge_status": "source_code_review",
                            "source_code_review_reason": (
                                "January source code differs from February onward; "
                                "Order440 historical code caption is Bezhеtsk; no code repair"
                            ),
                        }
                    )
                elif len(temporal_link) == 1:
                    reviewed = temporal_link.iloc[0]
                    record.update(
                        {
                            "bridge_status": "temporal_review",
                            "temporal_review_reason": reviewed.bridge_reason,
                            "candidate_territory_id": reviewed.territory_id,
                        }
                    )
                elif len(link) == 1:
                    linked = link.iloc[0]
                    record.update(
                        {
                            "territory_id": linked.territory_id,
                            "source_city_oktmo_11": linked.source_city_oktmo_11,
                            "in_primary_cohort": bool(linked.in_primary_cohort),
                            "bridge_status": "approved_city_context_only",
                        }
                    )
                cities.append(record)
                for col, (product, raw_value) in enumerate(zip(codes, row[2:], strict=True), 3):
                    value, status, literal = source_value(raw_value)
                    chunk.append(
                        (year, month, source_code, product, value, status, literal, row_number, col)
                    )
            if len(city_codes) != len(official):
                raise ValueError(
                    f"Incomplete city set {year}-{month}: {len(city_codes)} / {len(official)}"
                )
            frame = pd.DataFrame(chunk)
            # Explicit ordered insert avoids relying on inferred string column names.
            con.register("chunk", frame)
            con.execute("INSERT INTO cells SELECT * FROM chunk")
            con.unregister("chunk")
            qa.append(
                {
                    "year": year,
                    "month": month,
                    "source_city_rows": len(city_codes),
                    "product_columns": len(codes),
                    "source_cells": len(chunk),
                    "other_coded_geography_rows": noncity_rows,
                }
            )
            print(f"Extracted {year}-{month:02}: {len(chunk)} city-price cells", flush=True)
    source_output = args.output / "city_monthly_price_cells.parquet"
    con.execute(f"COPY cells TO '{source_output.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)")
    city_frame = pd.DataFrame(cities)
    con.register("city_dimensions", city_frame)
    primary = args.output / "primary_city_price_context.parquet"
    con.execute(f"""COPY (
        SELECT c.*, d.territory_id, d.source_city_oktmo_11, 'city_context_only' AS context_scope
        FROM cells c JOIN city_dimensions d USING (year, month, source_city_code_literal)
        WHERE d.in_primary_cohort AND d.bridge_status='approved_city_context_only'
        ) TO '{primary.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)""")
    summary = {
        "source_cell_count": con.sql("SELECT count(*) FROM cells").fetchone()[0],
        "numeric_count": con.sql("SELECT count(value) FROM cells").fetchone()[0],
        "source_status_counts": dict(
            con.sql("SELECT value_status,count(*) FROM cells GROUP BY 1").fetchall()
        ),
        "primary_context_cells": con.sql(
            f"SELECT count(*) FROM read_parquet('{primary.as_posix()}')"
        ).fetchone()[0],
        "primary_context_numeric": con.sql(
            f"SELECT count(value) FROM read_parquet('{primary.as_posix()}')"
        ).fetchone()[0],
        "source_grain": (
            "year/month/source city code/product code; end-of-period RUB per labelled unit"
        ),
        "no_imputation": True,
        "not_official_cpi": True,
    }
    summary["nominal_grid_cells"] = sum(
        len(lists.loc[lists.year == year])
        * 12
        * len({r["product_code"] for r in product_rows if r["year"] == year})
        for year in (2023, 2024)
    )
    summary["absent_product_column_cells"] = (
        summary["nominal_grid_cells"] - summary["source_cell_count"]
    )
    summary["blank_cells"] = summary["source_status_counts"].get("source_blank", 0)
    con.register("city_dimensions", city_frame)
    summary["source_code_review_cells"] = con.sql("""SELECT count(*) FROM cells c
        JOIN city_dimensions d USING (year,month,source_city_code_literal)
        WHERE d.bridge_status='source_code_review'""").fetchone()[0]
    summary["temporal_review_source_cells"] = con.sql("""SELECT count(*) FROM cells c
        JOIN city_dimensions d USING (year,month,source_city_code_literal)
        WHERE d.bridge_status='temporal_review'""").fetchone()[0]
    pd.DataFrame(product_rows).to_csv(args.output / "product_dictionary.csv", index=False)
    city_frame.to_csv(args.output / "city_monthly_dimensions.csv", index=False)
    pd.DataFrame(qa).to_csv(args.output / "monthly_extraction_qa.csv", index=False)
    pd.DataFrame(source_files).to_csv(args.output / "source_workbooks.csv", index=False)
    yearly_coverage = con.sql("""SELECT year, count(*) AS present_published_cells,
               count(value) AS numeric_cells,
               count(*) FILTER (WHERE value=0) AS source_zero_cells,
               count(*) FILTER (WHERE value_status='source_marker') AS marker_cells,
               count(*) FILTER (WHERE value_status='source_blank') AS blank_cells,
               count(distinct source_city_code_literal) AS distinct_source_codes
               FROM cells GROUP BY year ORDER BY year""").df()
    yearly_coverage["distinct_source_city_labels"] = yearly_coverage.year.map(
        city_frame.groupby("year").city_name_key.nunique()
    )
    yearly_coverage["nominal_grid_cells"] = yearly_coverage.year.map(
        {
            year: len(lists.loc[lists.year == year])
            * 12
            * len({r["product_code"] for r in product_rows if r["year"] == year})
            for year in (2023, 2024)
        }
    )
    yearly_coverage["absent_product_column_cells"] = (
        yearly_coverage.nominal_grid_cells - yearly_coverage.present_published_cells
    )
    yearly_coverage.to_csv(args.output / "yearly_source_coverage.csv", index=False)
    con.sql("""SELECT year, source_city_code_literal, count(*) AS expected_published_cells,
               count(value) AS numeric_cells, count(*) FILTER (WHERE value=0) AS source_zero_cells
               FROM cells GROUP BY 1,2 ORDER BY 1,2""").df().to_csv(
        args.output / "city_year_coverage.csv", index=False
    )
    con.sql(f"""SELECT year, product_code, count(*) AS expected_primary_cells,
               count(value) AS numeric_primary_cells,
               count(distinct territory_id) FILTER (WHERE value>0) AS municipalities_positive,
               count(distinct month) AS published_months,
               count(*) FILTER (WHERE value=0) AS source_zero_cells
               FROM read_parquet('{primary.as_posix()}') GROUP BY 1,2 ORDER BY 1,2""").df().to_csv(
        args.output / "product_primary_coverage.csv", index=False
    )
    for filename, obj in [("summary.json", summary), ("source_footnotes.json", footnotes)]:
        (args.output / filename).write_text(
            json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf8"
        )
    manifest = {
        "script_sha256": sha(Path(__file__)),
        "bridge_manifest_sha256": sha(args.bridge_dir / "manifest.json"),
        "inputs": [
            {"path": str(p), "sha256": sha(p)}
            for p in [lists_path, approved_path, candidates_path, annual_path]
        ],
        "source_workbooks": source_files,
        "outputs": [
            {"path": p.name, "sha256": sha(p), "bytes": p.stat().st_size}
            for p in sorted(args.output.iterdir())
        ],
    }
    (args.output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
