"""Download selected census evidence and expose source cells without joining geography."""

import argparse
import hashlib
import json
import re
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit, urlunsplit
from xml.etree import ElementTree

import duckdb
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from sber_enrichment.clients import PublicFetcher


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sources", type=Path, required=True)
    parser.add_argument("--raw", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--ca-bundle", type=Path)
    parser.add_argument("--offline", action="store_true")
    args = parser.parse_args()
    if args.out.exists() and any(args.out.iterdir()):
        raise FileExistsError("Use a new output directory")
    args.raw.mkdir(parents=True, exist_ok=True)
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "source_text").mkdir(exist_ok=True)
    sources = json.loads(args.sources.read_text(encoding="utf-8"))
    records, sheet_records, cells, hits = [], [], [], []
    for spec in sources:
        source_id = spec["source_id"]
        suffix = Path(unquote(urlsplit(spec["url"]).path)).suffix.lower()
        file = args.raw / (source_id + suffix)
        meta_file = file.with_suffix(suffix + ".json")
        if not file.exists():
            if args.offline:
                raise FileNotFoundError(file.name)
            host = urlsplit(spec["url"]).hostname
            ca = {host: args.ca_bundle} if args.ca_bundle else None
            parts = urlsplit(spec["url"])
            encoded_url = urlunsplit(parts._replace(path=quote(parts.path, safe="/%:@-._~")))
            fetched = PublicFetcher(timeout=40, max_bytes=30_000_000, ca_bundles=ca).fetch(
                encoded_url
            )
            body = fetched.pop("body")
            fetched["retrieved_at_utc"] = datetime.now(UTC).isoformat()
            write_json(meta_file, fetched)
            if fetched["error"]:
                records.append({**spec, "status": "fetch_failed", "error": fetched["error"]})
                continue
            file.write_bytes(body)
            time.sleep(1)
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
        assert sha(file) == meta["sha256"], file.name
        record = {
            **spec,
            "file": file.name,
            "sha256": sha(file),
            "bytes": file.stat().st_size,
            "retrieved_at_utc": meta["retrieved_at_utc"],
            "status": "downloaded",
        }
        if suffix == ".xlsx":
            book = load_workbook(file, read_only=False, data_only=False)
            record["sheet_count"] = len(book.sheetnames)
            for sheet in book:
                filled, numeric, formulas = 0, 0, 0
                header = []
                for row in sheet:
                    for cell in row:
                        value = cell.value
                        if value is None:
                            continue
                        filled += 1
                        is_number = isinstance(value, (int, float)) and not isinstance(value, bool)
                        numeric += int(is_number)
                        formulas += int(cell.data_type == "f")
                        raw = str(value)
                        cells.append(
                            {
                                "source_id": source_id,
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "row": cell.row,
                                "column": cell.column,
                                "raw_value": raw,
                                "numeric_value": float(value) if is_number else None,
                                "cell_type": cell.data_type,
                                "source_sha256": record["sha256"],
                            }
                        )
                        if cell.row <= 16:
                            header.append(f"{cell.coordinate}: {raw}")
                        if re.search(
                            "Нижнеколым|Верхнеколым|Свободненск|Куркино|Красноярск|Светловск",
                            raw,
                            re.I,
                        ):
                            hits.append(
                                {
                                    "source_id": source_id,
                                    "sheet": sheet.title,
                                    "cell": cell.coordinate,
                                    "territory_label": raw,
                                    "status": "text_locator_only_not_geographic_join",
                                }
                            )
                sheet_records.append(
                    {
                        "source_id": source_id,
                        "sheet": sheet.title,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                        "nonempty_cells": filled,
                        "numeric_cells": numeric,
                        "formula_cells": formulas,
                        "merged_ranges": ";".join(str(x) for x in sheet.merged_cells.ranges),
                        "header_excerpt": "\n".join(header),
                    }
                )
            book.close()
        elif suffix == ".pdf":
            pdf = PdfReader(file)
            text = "\n\n".join(f"PAGE {i + 1}\n{p.extract_text()}" for i, p in enumerate(pdf.pages))
            (args.out / "source_text" / (source_id + ".txt")).write_text(text, encoding="utf-8")
        elif suffix == ".docx":
            with zipfile.ZipFile(file) as z:
                root = ElementTree.fromstring(z.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            text = "\n".join("".join(p.itertext()) for p in root.findall(".//w:p", ns))
            (args.out / "source_text" / (source_id + ".txt")).write_text(text, encoding="utf-8")
        records.append(record)
        print(
            json.dumps(
                {"source_id": source_id, "status": record["status"], "bytes": record["bytes"]}
            ),
            flush=True,
        )
    pd.DataFrame(records).to_csv(args.out / "source_inventory.csv", index=False)
    pd.DataFrame(sheet_records).to_csv(args.out / "sheet_inventory.csv", index=False)
    pd.DataFrame(hits).to_csv(args.out / "six_case_text_locators.csv", index=False)
    frame = pd.DataFrame(cells)
    assert not frame.duplicated(["source_id", "sheet", "cell"]).any()
    duckdb.from_df(frame).write_parquet(str(args.out / "source_cells.parquet"), compression="zstd")
    manifest = {
        "status": "source_cells_ready_geographic_join_unresolved",
        "sources": len(records),
        "downloaded": sum(r["status"] == "downloaded" for r in records),
        "nonempty_cells": len(frame),
        "numeric_cells": int(frame.numeric_value.notna().sum()),
        "sheets": len(sheet_records),
        "six_case_text_locators": len(hits),
        "census_reference_date": "2021-10-01",
        "municipal_2023_2024_joins_accepted": 0,
        "sources_config_sha256": sha(args.sources),
        "script_sha256": sha(Path(__file__)),
        "source_hashes": {r["file"]: r["sha256"] for r in records if "file" in r},
        "outputs": {
            p.relative_to(args.out).as_posix(): sha(p) for p in args.out.rglob("*") if p.is_file()
        },
        "limits": [
            "No imputation of blanks or dash markers",
            "No automatic geographic joins",
            "Commuting outside settlement does not imply crossing municipality boundary",
            "Census 2021 values are not observations of 2023 or 2024",
            "Raw numeric cells include headers and totals; not a count of analytical observations",
        ],
    }
    write_json(args.out / "manifest.json", manifest)
    print(
        json.dumps(
            {k: v for k, v in manifest.items() if k not in ("source_hashes", "outputs")},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
