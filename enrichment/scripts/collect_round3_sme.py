"""Fetch a reviewed, size-bounded public SME source plan with byte-level provenance.

This collector neither joins municipalities nor treats a current extract as historical.
Each plan item declares its observation period, geography, and unresolved join limitations.
No credentials, model services, or retries are used.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import ssl
import time
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import urlsplit

import httpx
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def valid_official_url(url: str) -> bool:
    parsed = urlsplit(url)
    host = parsed.hostname or ""
    return (
        parsed.scheme == "https"
        and parsed.port in (None, 443)
        and not parsed.username
        and not parsed.password
        and any(
            host == base or host.endswith("." + base)
            for base in ("nalog.ru", "nalog.gov.ru", "rosstat.gov.ru")
        )
    )


def digest(body: bytes) -> str:
    return hashlib.sha256(body).hexdigest()


def write_json(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", "utf-8")


def fetch_one(client: httpx.Client, item: dict, root: Path, default_max_bytes: int) -> dict:
    source_id = item["source_id"]
    if not source_id.replace("-", "").replace("_", "").isalnum():
        raise ValueError("Unsafe source_id")
    if not valid_official_url(item["url"]):
        raise ValueError("Expected an official HTTPS URL")
    extension = item.get("extension", "raw")
    if not extension.isalnum():
        raise ValueError("Unsafe extension")
    method = item.get("method", "GET").upper()
    if method not in {"GET", "POST", "HEAD"}:
        raise ValueError("Only read-oriented GET/POST/HEAD plans are supported")
    meta_path = root / f"{source_id}.metadata.json"
    body_path = root / f"{source_id}.{extension}"
    if meta_path.exists():
        previous = json.loads(meta_path.read_text("utf-8"))
        if previous["request"] != item:
            raise ValueError(f"Changed existing source plan: {source_id}")
        if previous.get("sha256"):
            if digest(body_path.read_bytes()) != previous["sha256"]:
                raise ValueError(f"Cached source hash mismatch: {source_id}")
        return previous
    meta = {"request": item, "captured_at_utc": datetime.now(UTC).isoformat()}
    limit = int(item.get("max_bytes", default_max_bytes))
    try:
        with client.stream(method, item["url"], data=item.get("data")) as response:
            if not valid_official_url(str(response.url)):
                raise ValueError("Redirect left official allowed domains")
            meta.update(
                {
                    "http_status": response.status_code,
                    "final_url": str(response.url),
                    "response_headers": {
                        key: value
                        for key, value in response.headers.items()
                        if key.lower()
                        in {
                            "content-type",
                            "content-length",
                            "last-modified",
                            "etag",
                            "content-disposition",
                            "accept-ranges",
                            "content-range",
                        }
                    },
                }
            )
            announced = response.headers.get("content-length")
            meta["announced_bytes"] = int(announced) if announced and announced.isdigit() else None
            if method == "HEAD":
                meta["access_status"] = "headers_only_size_assessment"
            elif meta["announced_bytes"] is not None and meta["announced_bytes"] > limit:
                meta["access_status"] = "not_downloaded_exceeds_byte_limit"
            else:
                chunks = []
                received = 0
                for chunk in response.iter_bytes():
                    received += len(chunk)
                    if received > limit:
                        raise ValueError("Stream exceeded byte limit; no partial body retained")
                    chunks.append(chunk)
                body = b"".join(chunks)
                body_path.write_bytes(body)
                meta.update(
                    {
                        "bytes": len(body),
                        "sha256": digest(body),
                        "path": body_path.name,
                        "access_status": "downloaded" if response.is_success else "http_error",
                    }
                )
    except (httpx.HTTPError, ValueError) as exc:
        meta.update({"access_status": "fetch_failed", "error": f"{type(exc).__name__}: {exc}"})
    write_json(meta_path, meta)
    return meta


def numeric_cell(value):
    """Source markers remain missing; neither dash nor blank is imputed as zero."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value, "observed"
    return None, "source_blank" if value is None else "unresolved_source_marker"


def vologda_counts(path: Path, year: int, source: dict) -> list[dict]:
    """Extract source municipality/industry counts without any external name joins."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook["T25"].iter_rows(values_only=True))
    title = str(rows[0][0])
    if f"{year} году" not in title or "Группировка малых предприятий" not in title:
        raise ValueError("Municipal workbook title/year/scope changed")
    if "единиц" not in str(rows[2][1]) or "количе" not in str(rows[3][3]).lower():
        raise ValueError("Organization count units/column positions changed")
    municipality = None
    seen = set()
    result = []
    for offset, cells in enumerate(rows):
        label = str(cells[0]).strip() if cells[0] is not None else ""
        if label == "ВСЕГО":
            previous = rows[offset - 1]
            if any(value is not None for value in previous[1:]) or not isinstance(previous[0], str):
                raise ValueError("Municipal heading is not an explicit separate row")
            municipality = previous[0]  # Includes footnote3 where present, without silent repair.
            if municipality in seen:
                raise ValueError("Duplicate municipality block")
            seen.add(municipality)
        if municipality is None or not label or cells[1] is None:
            continue
        if label.startswith(("1 Информация", "2 По сопоставимому", "3 Включая")):
            break
        count, status = numeric_cell(cells[1])
        if count is not None and (count < 0 or count != int(count)):
            raise ValueError("Organization counts must be nonnegative integers")
        profit, profit_status = numeric_cell(cells[3])
        loss, loss_status = numeric_cell(cells[6])
        if all(value is not None for value in (count, profit, loss)) and count != profit + loss:
            raise ValueError("Published organization grouping does not reconcile")
        result.append(
            {
                "source_id": source["request"]["source_id"],
                "observation_year": year,
                "source_municipality_label": municipality,
                "source_activity_label": label,
                "is_total": label == "ВСЕГО",
                "organizations": count,
                "status": status,
                "organizations_raw": cells[1],
                "profitable_organizations": profit,
                "profitable_raw": cells[3],
                "profitable_status": profit_status,
                "loss_organizations": loss,
                "loss_raw": cells[6],
                "loss_status": loss_status,
                "unit": "organizations",
                "scope": "small_enterprises_in_financial_source_population",
                "includes_micro_explicitly_confirmed": None,
                "source_sheet": "T25",
                "source_count_cell": f"B{offset + 1}",
                "source_row": offset + 1,
                "source_sha256": source["sha256"],
                "source_url": source["request"]["url"],
                "territory_id": None,
                "official_oktmo_8": None,
                "join_status": "unlinked_no_municipal_code_in_source",
                "boundary_equivalence_confirmed": None,
            }
        )
    return result


def rosstat_regional_workers(path: Path, source: dict) -> list[dict]:
    """Keep explicit years/cells and nesting labels; never infer a municipal allocation."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = []
    for sheet_name, geography in [
        ("2", "Russian_Federation_by_activity"),
        ("4", "published_region_or_parent_aggregate"),
    ]:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        title = str(rows[1][0])
        if "Средняя численность" not in title or "человек" not in title:
            raise ValueError("Worker concept/units changed")
        year_columns = {
            int(value): col
            for col, value in enumerate(rows[3])
            if isinstance(value, (int, float)) and value in (2023, 2024)
        }
        if set(year_columns) != {2023, 2024}:
            raise ValueError("Explicit requested annual headers absent")
        for offset, cells in enumerate(rows[4:], 4):
            if cells[0] is None or all(cells[col] is None for col in year_columns.values()):
                continue
            for year, col in year_columns.items():
                value, status = numeric_cell(cells[col])
                result.append(
                    {
                        "source_id": source["request"]["source_id"],
                        "observation_year": year,
                        "geography_grain": geography,
                        "source_label": cells[0],
                        "value": value,
                        "value_raw": cells[col],
                        "status": status,
                        "unit": "people",
                        "concept": "average_number_of_workers_not_average_payroll_only",
                        "enterprise_scope": (
                            "small_including_micro"
                            if "включая микропредприятия" in title
                            else "small_excluding_micro"
                        ),
                        "source_sheet": sheet_name,
                        "source_cell": f"{get_column_letter(col + 1)}{offset + 1}",
                        "source_sha256": source["sha256"],
                        "source_url": source["request"]["url"],
                        "territory_id": None,
                        "join_status": "not_municipal_no_allocation",
                        "not_to_sum_rows": True,
                    }
                )
    return result


def fns_statistics(path: Path, source: dict) -> list[dict]:
    rows = json.loads(path.read_text("utf-8"))
    requested_date = source["request"]["data"]["statDate"]
    parsed_date = datetime.strptime(requested_date, "%d.%m.%Y")
    worker_year = parsed_date.year - (2 if parsed_date.month <= 6 else 1)
    result = []
    for index, row in enumerate(rows):
        if row["stat_date"].split(" ")[0] != requested_date:
            raise ValueError("FNS ignored historical snapshot date")
        if row["cnt_total"] != row["cnt_ul_total"] + row["cnt_ip_total"]:
            raise ValueError("FNS entity total does not reconcile")
        if row["cnt_worker_total"] != row["cnt_worker_ul_total"] + row["cnt_worker_ip_total"]:
            raise ValueError("FNS worker total does not reconcile")
        record = {
            "source_id": source["request"]["source_id"],
            "snapshot_date": parsed_date.date().isoformat(),
            "employee_reference_year": worker_year,
            "entity_count_is_snapshot_not_annual_flow": True,
            "source_level": row["cnt_level"],
            "source_region_code_ssrf": row.get("ssrf"),
            "source_federal_district_code": row.get("focode"),
            "source_name": row["cnt_name"],
            "source_sha256": source["sha256"],
            "source_url": source["request"]["url"],
            "source_json_array_index": index,
            "territory_id": None,
            "join_status": "regional_FNS_codes_no_municipal_join",
        }
        # Keep original count fields, including micro/small/medium and legal entities/IE.
        record.update(
            {
                key: value
                for key, value in row.items()
                if key.startswith("cnt_") and key not in {"cnt_name", "cnt_level", "cnt_is_bold"}
            }
        )
        result.append(record)
    return result


def derive(raw: Path, out: Path) -> dict:
    if out.exists() and any(out.iterdir()):
        raise ValueError("Derived output must be new or empty")
    inventory = json.loads((raw / "inventory.json").read_text("utf-8"))
    by_id = {row["request"]["source_id"]: row for row in inventory}
    for record in inventory:
        if record.get("sha256") and digest((raw / record["path"]).read_bytes()) != record["sha256"]:
            raise ValueError("Input hash mismatch")
    municipal = []
    for source_id, year in [
        ("vologda-finance-page-table-4", 2023),
        ("vologda-finance-page-table-2", 2024),
    ]:
        source = by_id[source_id]
        municipal.extend(vologda_counts(raw / source["path"], year, source))
    regional = []
    for number in (1, 2):
        source = by_id[f"rosstat-sme-regional-page-table-{number}"]
        regional.extend(rosstat_regional_workers(raw / source["path"], source))
    fns = []
    for year in (2023, 2024, 2025):
        source = by_id[f"fns-statistics-{year}0710"]
        fns.extend(fns_statistics(raw / source["path"], source))
    municipal = pd.DataFrame(municipal)
    if municipal.duplicated(["source_id", "source_row"]).any():
        raise ValueError("Repeated municipal source row")
    if not municipal.loc[municipal.is_total].groupby("observation_year").size().eq(28).all():
        raise ValueError("Expected 28 explicit municipal blocks per year")
    regional = pd.DataFrame(regional)
    fns = pd.DataFrame(fns)
    out.mkdir(parents=True, exist_ok=True)
    municipal.to_csv(out / "vologda_small_organizations_unlinked.csv", index=False)
    regional.to_csv(out / "rosstat_regional_workers_context.csv", index=False)
    fns.to_csv(out / "fns_historical_aggregates.csv", index=False)
    fields = ("ОКТМО", "ОКАТО", "КЛАДР", "ФИАС", "КодРегион", "ССЧР")
    schema = {
        year: {
            field: field in (raw / by_id[f"fns-schema-{year}"]["path"]).read_text("utf-8")
            for field in fields
        }
        for year in (2022, 2023, 2025)
    }
    qa = {
        "readiness": "source_context_only_no_new_municipal_worker_join",
        "municipal_source_rows": len(municipal),
        "municipal_total_blocks_by_year": municipal.loc[municipal.is_total]
        .groupby("observation_year")
        .size()
        .to_dict(),
        "municipal_distinct_industry_labels_excluding_total": int(
            municipal.loc[~municipal.is_total, "source_activity_label"].nunique()
        ),
        "municipal_organization_count_numeric_cells": int(municipal.organizations.notna().sum()),
        "municipal_count_cell_statuses": municipal.status.value_counts().to_dict(),
        "fns_rows": len(fns),
        "fns_level_counts": fns.groupby("snapshot_date")
        .source_level.value_counts()
        .rename("n")
        .reset_index()
        .to_dict("records"),
        "fns_snapshot_and_reference_years": fns[["snapshot_date", "employee_reference_year"]]
        .drop_duplicates()
        .to_dict("records"),
        "rosstat_regional_context_rows": len(regional),
        "rosstat_regional_context_statuses": regional.status.value_counts().to_dict(),
        "fns_historical_schema_literal_field_audit": schema,
        "municipal_joins_made": 0,
        "municipal_workers2023_2024_gap_closed": False,
        "all_downloaded_source_hashes_verified": True,
        "archived_releases_not_downloaded": [
            {
                "source_id": row["request"]["source_id"],
                "bytes": row.get("announced_bytes"),
                "url": row["request"]["url"],
            }
            for row in inventory
            if row["request"].get("method") == "HEAD"
        ],
        "captured_sources": len(inventory),
        "downloaded_bytes": sum(row.get("bytes", 0) for row in inventory),
        "medium_scope_overlap_prevents_naive_addition": True,
        "current_snapshots_not_substituted_for_historical": True,
        "missing_and_source_markers_not_zero_imputed": True,
    }
    write_json(out / "QA.json", qa)
    write_json(out / "source_inventory.json", inventory)
    write_json(
        out / "manifest.json",
        {
            "created_at_utc": datetime.now(UTC).isoformat(),
            "script_sha256": digest(Path(__file__).read_bytes()),
            "source_inventory_sha256": digest((raw / "inventory.json").read_bytes()),
            "runtime_lock_sha256": digest((Path(__file__).parents[1] / "uv.lock").read_bytes()),
            "output_hashes": {
                path.name: digest(path.read_bytes()) for path in sorted(out.iterdir())
            },
            "input_paths_relative_to_raw_dir": {
                row["path"]: row["sha256"] for row in inventory if row.get("sha256")
            },
        },
    )
    return qa


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--plan", type=Path, required=True)
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--max-bytes", type=int, default=25_000_000)
    parser.add_argument("--interval", type=float, default=1.5)
    parser.add_argument("--derive-only", action="store_true")
    parser.add_argument("--report-dir", type=Path)
    args = parser.parse_args()
    items = json.loads(args.plan.read_text("utf-8"))
    if len({item["source_id"] for item in items}) != len(items):
        raise ValueError("Duplicate source identifiers")
    if args.derive_only:
        if args.report_dir is None:
            raise ValueError("--derive-only requires --report-dir")
        print(json.dumps(derive(args.raw_dir, args.report_dir), ensure_ascii=False))
        return
    args.raw_dir.mkdir(parents=True, exist_ok=True)
    results = []
    with httpx.Client(
        verify=ssl.create_default_context(), timeout=45, follow_redirects=True
    ) as client:
        for item in items:
            result = fetch_one(client, item, args.raw_dir, args.max_bytes)
            results.append(result)
            print(
                json.dumps(
                    {
                        "source_id": item["source_id"],
                        "status": result["access_status"],
                        "bytes": result.get("bytes"),
                        "announced_bytes": result.get("announced_bytes"),
                    }
                ),
                flush=True,
            )
            write_json(args.raw_dir / "inventory.json", results)
            time.sleep(max(0, args.interval))
    if args.report_dir:
        print(json.dumps(derive(args.raw_dir, args.report_dir), ensure_ascii=False))


if __name__ == "__main__":
    main()
