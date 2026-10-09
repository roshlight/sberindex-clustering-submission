"""Extract dated official price context, preserving source cells and territorial levels."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def norm(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower().replace("ё", "е")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def prepare(raw, repo, out):
    out.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((raw / "download_manifest.json").read_text(encoding="utf-8"))
    sources = {r["filename"]: r for r in manifest if r.get("status") == 200}
    for name, record in sources.items():
        if digest(raw / name) != record["sha256"]:
            raise ValueError(f"Source changed: {name}")

    def provenance(name, sheet, row, column):
        source = sources[name]
        return {
            "source_file": name,
            "source_sheet": sheet,
            "source_cell": f"{get_column_letter(column)}{row}",
            "source_url": source["url"],
            "source_sha256": source["sha256"],
            "retrieved_at": source["retrieved_at"],
            "published_at": None,
        }

    name = "ipc_RF_fo_sub_08-2026.xlsx"
    wb = load_workbook(raw / name, read_only=True, data_only=True)
    rows = []
    labels = {3: "all", 4: "food", 5: "nonfood", 6: "services"}
    for year in (2023, 2024):
        for month in range(1, 13):
            sheet = wb[f"{month:02d}({year})"]
            if "к предыдущему месяцу" not in str(sheet["C2"].value):
                raise ValueError("Unexpected CPI base")
            if norm(sheet["A3"].value) != "код территории":
                raise ValueError("Unexpected CPI code column")
            for number, record in enumerate(sheet.values, 1):
                code, title = record[:2]
                if not isinstance(code, int):
                    continue
                level = "country" if code == 643 else "federal_district" if code < 100 else "region"
                if code in (11000000, 71000000):
                    level = "composite_region_including_autonomous_okrugs"
                for col, category in labels.items():
                    value = record[col - 1]
                    if not isinstance(value, (int, float)) or not 50 < value < 200:
                        raise ValueError("Missing or implausible source CPI")
                    rows.append(
                        {
                            "source_territory_code": str(code),
                            "source_territory_name": title,
                            "territorial_level": level,
                            "month": f"{year}-{month:02d}",
                            "category": category,
                            "value": value,
                            "unit": "percent",
                            "index_base": "previous_month_100",
                            **provenance(name, sheet.title, number, col),
                        }
                    )
    cpi = pd.DataFrame(rows)
    key = ["source_territory_code", "month", "category"]
    if cpi.duplicated(key).any() or cpi.groupby(key[0]).size().nunique() != 1:
        raise ValueError("Duplicate or incomplete source CPI series")
    cpi.to_csv(out / "cpi_monthly_2023_2024.csv", index=False)

    # Labels are reconciled explicitly. Source price codes are not declared to be OKTMO.
    territories = cpi[["source_territory_code", "source_territory_name"]].drop_duplicates()
    lookup = {
        norm(r.source_territory_name): r.source_territory_code for r in territories.itertuples()
    }
    overrides = {
        "15": ("90000000", "Республика Северная Осетия - Алания"),
        "83": ("11100000", "Ненецкий авт.округ"),
        "42": ("32000000", "Кемеровская область - Кузбасс"),
        "86": ("71100000", "Ханты-Мансийский авт.округ - Югра"),
        "89": ("71140000", "Ямало-Hенецкий авт.округ"),
        "79": ("99000000", "Еврейская авт.область"),
        "87": ("77000000", "Чукотский авт.округ"),
        "78": ("40000000", "г.Санкт-Петербург"),
        "77": ("45000000", "г. Москва"),
        "92": ("67000000", "г.Севастополь"),
        "29": ("11001000", "Архангельская область (кроме Ненецкого автономного округа)"),
        "72": (
            "71001000",
            "Тюменская область (кроме Ханты-Мансийского автономного округа - Югры "
            "и Ямало-Ненецкого автономного округа)",
        ),
    }
    annual = []
    for year in (2023, 2024):
        cross = pd.read_csv(
            repo / f"reports/enrichment_v1/crosswalk/municipal_crosswalk_{year}-01-01.csv",
            dtype=str,
        )
        annual.append(cross[["territory_id", "snapshot_date", "region_code", "region_name"]])
    annual = pd.concat(annual, ignore_index=True)
    regions = annual[["region_code", "region_name"]].drop_duplicates()
    bridge = []
    for row in regions.itertuples():
        if row.region_code in overrides:
            code, expected_name = overrides[row.region_code]
            if lookup.get(norm(expected_name)) != code:
                raise ValueError(f"Regional alias no longer valid: {row.region_code}")
            rule = "explicit_label_review_excluding_nested_okrugs"
        else:
            code = lookup.get(norm(row.region_name))
            rule = "exact_region_label_after_whitespace_case_normalization"
        if code is None:
            raise ValueError(f"Unmatched region: {row.region_name}")
        source_name = territories.set_index("source_territory_code").loc[
            code, "source_territory_name"
        ]
        bridge.append(
            {
                "region_code": row.region_code,
                "region_name": row.region_name,
                "source_territory_code": code,
                "source_territory_name": source_name,
                "mapping_rule": rule,
            }
        )
    bridge = pd.DataFrame(bridge)
    if bridge.region_code.duplicated().any() or bridge.source_territory_code.duplicated().any():
        raise ValueError("Regional crosswalk is not one-to-one")
    bridge.to_csv(out / "region_price_bridge.csv", index=False)
    cohort = pd.read_csv(
        repo / "reports/enrichment_v2/population_source/target_reference_links.csv", dtype=str
    )
    cohort = cohort[["territory_id", "reference_date"]].drop_duplicates()
    links = cohort.merge(
        annual,
        left_on=["territory_id", "reference_date"],
        right_on=["territory_id", "snapshot_date"],
        validate="one_to_one",
    )
    links = links.merge(
        bridge[["region_code", "source_territory_code", "mapping_rule"]],
        on="region_code",
        validate="many_to_one",
    )
    if len(links) != 4032 or links.source_territory_code.isna().any():
        raise ValueError("Incomplete annual municipality-region links")
    links["data_level"] = "region; repeated municipal links are not independent price observations"
    links.to_csv(out / "municipal_region_price_links.csv", index=False)

    name = "isg_2012-2025.xlsx"
    sheet = load_workbook(raw / name, read_only=True, data_only=True)["1"]
    if [sheet.cell(3, c).value for c in (13, 14)] != [2023, 2024]:
        raise ValueError("Cost-of-living year headers changed")
    city_rows = []
    district = None
    for number, row in enumerate(sheet.values, 1):
        title = row[0]
        if isinstance(title, str) and "федеральный округ" in title:
            district = title
        if number < 4 or not isinstance(title, str):
            continue
        for year, col in ((2023, 13), (2024, 14)):
            value = row[col - 1]
            if any(isinstance(row[c - 1], (int, float)) or row[c - 1] == "…" for c in (13, 14)):
                city_rows.append(
                    {
                        "source_city_name": title,
                        "federal_district": district,
                        "year": year,
                        "value": value if isinstance(value, (int, float)) else None,
                        "source_raw_value": value,
                        "value_status": "observed"
                        if isinstance(value, (int, float))
                        else "source_missing",
                        "unit": "ratio",
                        "index_base": "country_average_1",
                        "temporal_basis": "annual_average",
                        "territorial_level": "country"
                        if title == "Российская Федерация"
                        else "surveyed_city",
                        "municipal_join_status": "not_joined_city_is_not_municipality",
                        **provenance(name, sheet.title, number, col),
                    }
                )
    cities = pd.DataFrame(city_rows)
    cities.to_csv(out / "city_cost_of_living_2023_2024.csv", index=False)

    name = "fixed_basket_a5413194.xlsx"
    wb = load_workbook(raw / name, read_only=True, data_only=True)
    baskets = []
    for month in range(1, 13):
        sheet = wb[str(month)]
        if "2024 года" not in str(sheet["A2"].value):
            raise ValueError("Fixed basket year is not 2024")
        for number, row in enumerate(sheet.values, 1):
            title, value = row[:2]
            if not isinstance(title, str) or not isinstance(value, (int, float)):
                continue
            if not 1000 < value < 100000:
                raise ValueError("Implausible fixed basket cost")
            baskets.append(
                {
                    "source_territory_name": title,
                    "month": f"2024-{month:02d}",
                    "value": value,
                    "unit": "rubles_per_fixed_basket_per_month",
                    "source_mom_percent": row[2],
                    "source_december_previous_year_percent": row[3],
                    "municipal_join_status": "not_joined_regional_context",
                    **provenance(name, sheet.title, number, 2),
                }
            )
    baskets = pd.DataFrame(baskets)
    basket_aliases = {
        "Город Москва столица Российской Федерации город федерального значения": "45000000",
        "Ненецкий автономный округ (Архангельская область)": "11100000",
        "Город Санкт-Петербург город федерального значения": "40000000",
        "Южный федеральный округ (с 29.07.2016)": "40",
        "Республика Адыгея (Адыгея)": "79000000",
        "Город федерального значения Севастополь": "67000000",
        "Республика Северная Осетия-Алания": "90000000",
        "Чеченская Республика*": "96000000",
        "Республика Татарстан (Татарстан)": "92000000",
        "Чувашская Республика - Чувашия": "97000000",
        "Ханты-Мансийский автономный округ - Югра (Тюменская область)": "71100000",
        "Ямало-Ненецкий автономный округ (Тюменская область)": "71140000",
        "Тюменская область (кроме Ханты-Мансийского автономного округа-Югры "
        "и Ямало-Ненецкого автономного округа)": "71001000",
        "Еврейская автономная область": "99000000",
        "Чукотский автономный округ": "77000000",
    }
    basket_label_lookup = {**lookup, **{norm(k): v for k, v in basket_aliases.items()}}
    baskets["source_territory_code"] = baskets.source_territory_name.map(
        lambda x: basket_label_lookup.get(norm(x))
    )
    if baskets.source_territory_code.isna().any():
        raise ValueError("Unreviewed fixed basket region label")
    price_levels = cpi.set_index("source_territory_code").territorial_level.to_dict()
    baskets["territorial_level"] = baskets.source_territory_code.map(price_levels)
    baskets["regional_code_mapping"] = "exact_label_or_explicit_review_to_cpi_source_code"
    if baskets.duplicated(["source_territory_code", "month"]).any():
        raise ValueError("Duplicate reconciled basket territory")
    if baskets.duplicated(["source_territory_name", "month"]).any():
        raise ValueError("Duplicate basket observations")
    baskets.to_csv(out / "fixed_basket_regional_2024.csv", index=False)
    annual_name = "ipc_s_1992-2025.xlsx"
    annual_sheet = load_workbook(raw / annual_name, read_only=True, data_only=True).active
    annual_values = {}
    for row_number, values in enumerate(annual_sheet.values, 1):
        if not isinstance(values[0], str):
            continue
        for year in (2023, 2024):
            column = year - 1990
            if not str(annual_sheet.cell(4, column).value).startswith(str(year)):
                raise ValueError("Annual CPI year column moved")
            value = values[column - 1]
            if isinstance(value, (int, float)):
                key = (norm(values[0]), year)
                if key in annual_values:
                    raise ValueError("Duplicate annual CPI territory label")
                annual_values[key] = (value, row_number, column)
    all_cpi = cpi[cpi.category.eq("all")].copy()
    all_cpi["year"] = all_cpi.month.str[:4].astype(int)
    chained = all_cpi.groupby(["source_territory_name", "year"]).value.prod() / 100**11
    comparisons = []
    for (title, year), chained_value in chained.items():
        if (norm(title), year) not in annual_values:
            continue
        annual_value, row_number, column = annual_values[(norm(title), year)]
        comparisons.append(
            {
                "name": title,
                "year": year,
                "chain": chained_value,
                "annual": annual_value,
                "difference_pp": chained_value - annual_value,
                "monthly_source_sha256": sources["ipc_RF_fo_sub_08-2026.xlsx"]["sha256"],
                **provenance(annual_name, annual_sheet.title, row_number, column),
            }
        )
    comparisons = pd.DataFrame(comparisons)
    comparisons.to_csv(out / "cpi_annual_reconciliation.csv", index=False)
    result = {
        "status": "ready_with_conditions",
        "cpi_rows": len(cpi),
        "cpi_source_territories": cpi.source_territory_code.nunique(),
        "cpi_months": cpi.month.nunique(),
        "cpi_categories": sorted(cpi.category.unique().tolist()),
        "cpi_territorial_level_rows": cpi.territorial_level.value_counts().to_dict(),
        "region_bridge_rows": len(bridge),
        "cohort_links": len(links),
        "cohort_territories": links.territory_id.nunique(),
        "cohort_regions": links.region_code.nunique(),
        "cost_of_living_rows": len(cities),
        "cost_of_living_city_rows": int(cities.territorial_level.eq("surveyed_city").sum()),
        "cost_of_living_observed_city_values": int(
            (cities.territorial_level.eq("surveyed_city") & cities.value.notna()).sum()
        ),
        "cost_of_living_missing_city_values": int(cities.value.isna().sum()),
        "cost_of_living_observed_cities_by_year": cities.loc[
            cities.territorial_level.eq("surveyed_city") & cities.value.notna()
        ]
        .groupby("year")
        .size()
        .to_dict(),
        "fixed_basket_rows": len(baskets),
        "fixed_basket_source_territories": baskets.source_territory_code.nunique(),
        "fixed_basket_source_labels": baskets.source_territory_name.nunique(),
        "fixed_basket_territorial_level_rows": baskets.territorial_level.value_counts().to_dict(),
        "cpi_annual_comparisons": len(comparisons),
        "cpi_annual_max_difference_pp": comparisons.difference_pp.abs().max(),
        "source_files": [
            {k: v for k, v in r.items() if k != "original_url"} for r in sources.values()
        ],
        "limitations": [
            "2026 source vintage is retrospective evidence, not data available during 2023/2024.",
            "CPI change is not a spatial price level; "
            "a common deflator cannot change normalized profiles.",
            "Municipal region links do not create municipal price measurements.",
            "The four broad CPI groups cannot be mapped directly to the five Sber categories.",
            "City cost of living and fixed baskets are extracted; "
            "they are not automatically joined to municipalities.",
            "Original frozen experiment and enrichment_v1/v2 remain unchanged.",
        ],
    }
    result["output_hashes"] = {p.name: digest(p) for p in sorted(out.glob("*.csv"))}
    (out / "manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                k: v
                for k, v in result.items()
                if k not in {"source_files", "output_hashes", "limitations"}
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw", type=Path, required=True)
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    prepare(args.raw, args.repo, args.out)
