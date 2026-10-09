import hashlib
import json
import runpy
from pathlib import Path

import duckdb
import pandas as pd
import pytest
from openpyxl import Workbook

from sber_enrichment.population import REGIONS_SHEET, SHEET, parse_population_workbook

URL = "https://rosstat.gov.ru/storage/mediabank/BUL_MO_2024.xlsx"


def _book(path, *, with_codes=True, year=2024, rows=None, dash_definition=True):
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    offset = 1 if with_codes else 0
    shift = 0 if with_codes else 1
    sheet.cell(2 + shift, offset + 1, "Численность постоянного населения")
    sheet.cell(3 + shift, offset + 1, f"по муниципальным образованиям на 1 января {year} года")
    sheet.cell(5 + shift, offset + 2, "Все")
    sheet.cell(6 + shift, offset + 2, "население")
    sheet.cell(7 + shift, offset + 2, "(человек)")
    sheet.cell(6 + shift, offset + 3, "городское")
    sheet.cell(6 + shift, offset + 4, "сельское")
    if with_codes:
        sheet.cell(6, 1, "Коды территорий")
        sheet.cell(7, 1, "ТЕРСОН-МО")
    if rows is None:
        rows = [
            ("0100000000", "Тестовая область", 100, 60, 40),
            ("0170100000", "Городской округ Тестовый", 50, 40, 10),
            ("017010000011001", "г Тестовый", 40, 40, 0),
        ]
    for number, (code, name, total, urban, rural) in enumerate(rows, 8 + shift):
        values = [code, name, total, urban, rural] if with_codes else [name, total, urban, rural]
        for column, value in enumerate(values, 1):
            sheet.cell(number, column, value)
    regions = book.create_sheet(REGIONS_SHEET)
    regions.append([None, "Название"])
    regions.append([1, "Тестовая область"])
    preface = book.create_sheet("Предисловие")
    if dash_definition:
        preface.cell(1, 1, "-       - явление отсутствует;")
    preface.cell(2, 1, "… - данных не имеется.")
    book.save(path)


def _parse(path):
    return parse_population_workbook(path, source_url=URL, retrieved_at="2026-09-14")


def test_reference_date_is_read_from_content_and_terson_is_not_oktmo(tmp_path):
    path = tmp_path / "misleading_1999.xlsx"
    _book(path)
    frame, audit = _parse(path)
    assert frame.reference_date.eq("2024-01-01").all()
    assert frame.source_code.iloc[0] == "0100000000"
    assert frame.source_code_system.eq("TERSON-MO").all()
    assert frame.oktmo.isna().all() and frame.territory_id.isna().all()
    assert frame.region_name_source.eq("Тестовая область").all()
    assert frame.row_kind.tolist() == [
        "region_header",
        "municipality_explicit_label",
        "locality_explicit_label",
    ]
    assert audit["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert not frame.source_record_id.duplicated().any()
    assert audit["published_at"] is None


def test_old_workbook_keeps_uncoded_rows_unlinked(tmp_path):
    path = tmp_path / "old.xlsx"
    _book(path, with_codes=False, year=2023)
    frame, audit = _parse(path)
    assert frame.reference_date.eq("2023-01-01").all()
    assert frame.source_code.isna().all()
    assert frame.source_code_status.eq("not_provided").all()
    assert audit["source_code_system"] == "not_provided"
    assert frame.source_row.tolist() == [9, 10, 11]


def test_blank_ellipsis_explicit_absence_and_bad_balance_are_distinct(tmp_path):
    path = tmp_path / "markers.xlsx"
    _book(
        path,
        rows=[
            ("0100000000", "Тестовая область", 100, 60, 40),
            ("1", "Первое поселение", 10, None, 10),
            ("2", "Второе поселение", 10, "…", 10),
            ("3", "Третье поселение", 10, "-", 10),
            ("4", "Четвертое поселение", 10, 11, 0),
        ],
    )
    frame, audit = _parse(path)
    assert pd.isna(frame.population_urban.iloc[1])
    assert pd.isna(frame.population_urban.iloc[2])
    assert frame.population_urban.iloc[3] == 0
    assert frame.population_urban_status.iloc[3] == "source_absence_marker"
    assert frame.population_urban_raw.iloc[3] == "-"
    assert frame.balance_status.iloc[4] == "mismatch"
    assert pd.isna(frame.urban_share.iloc[4])
    assert audit["balance_mismatch_rows"] == [12]


def test_dash_requires_definition_and_codes_are_not_silently_repaired(tmp_path):
    path = tmp_path / "badcodes.xlsx"
    _book(
        path,
        dash_definition=False,
        rows=[
            ("0100000000", "Тестовая область", 100, 60, 40),
            ("01 70100000", "Первое поселение", 10, "-", 10),
            (1234567890123456, "Второе поселение", 10, 0, 10),
            ("same", "Третье поселение", 10, 0, 10),
            (1234, "Четвертое поселение", 10, 0, 10),
        ],
    )
    frame, _ = _parse(path)
    assert frame.source_code_status.iloc[1] == "unparsed_text"
    assert pd.isna(frame.population_urban.iloc[1])
    assert frame.source_code_status.iloc[2] == "numeric_precision_risk"
    assert pd.isna(frame.source_code.iloc[2])
    assert frame.source_code_status.iloc[4] == "numeric_leading_zeros_unverified"


def test_duplicate_code_rows_are_preserved_and_flagged(tmp_path):
    path = tmp_path / "duplicates.xlsx"
    _book(
        path,
        rows=[
            ("0100000000", "Тестовая область", 100, 60, 40),
            ("2", "Первое поселение", 10, 0, 10),
            ("2", "Второе поселение", 20, 10, 10),
        ],
    )
    frame, audit = _parse(path)
    assert len(frame) == 3
    assert audit["duplicate_source_code_rows"] == 2
    assert frame.source_code_duplicate.sum() == 2


def test_units_must_be_explicit(tmp_path):
    path = tmp_path / "wrongunits.xlsx"
    _book(path)
    from openpyxl import load_workbook

    book = load_workbook(path)
    book[SHEET]["C7"] = "тыс. человек"
    book.save(path)
    with pytest.raises(ValueError, match="units"):
        _parse(path)


def test_uncalculated_formula_is_not_silently_skipped(tmp_path):
    path = tmp_path / "formula.xlsx"
    _book(path, rows=[("0100000000", "Тестовая область", "=60+40", "=60", "=40")])
    with pytest.raises(ValueError, match="formula"):
        _parse(path)


def test_output_artifacts_keep_missing_oktmo_and_refuse_overwrite(tmp_path):
    path = tmp_path / "input.xlsx"
    _book(path)
    prepare = runpy.run_path(
        str(Path(__file__).resolve().parents[1] / "scripts" / "prepare_population.py")
    )["prepare"]
    output = tmp_path / "parsed"
    manifest = prepare([path], [URL], output, "2026-09-14")
    assert manifest["status"] == "complete_unlinked"
    with duckdb.connect() as connection:
        frame = connection.read_parquet(str(output / "population_unlinked.parquet")).df()
    assert frame.oktmo.isna().all()
    for filename, digest in manifest["output_hashes"].items():
        assert hashlib.sha256((output / filename).read_bytes()).hexdigest() == digest
    assert (
        json.loads((output / "population_source_audit.json").read_text(encoding="utf-8"))[0][
            "n_rows"
        ]
        == 3
    )
    with pytest.raises(FileExistsError):
        prepare([path], [URL], output, "2026-09-14")
