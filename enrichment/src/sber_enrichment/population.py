"""Extract Rosstat population bulletins without inventing a municipal crosswalk.

TЕРСОН-МО is kept as the source's code system, never relabelled as OKTMO.
Physical worksheet rows are authoritative keys because source codes may repeat.
"""

from __future__ import annotations

import hashlib
import math
import re
from collections import Counter
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SHEET = "Численность_по_МО"
REGIONS_SHEET = "Перечень_субъектов_РФ"
POPULATION_FIELDS = ("population_total", "population_urban", "population_rural")


def _text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _population_value(value, *, absence_marker_defined: bool):
    if value is None or _text(value) == "":
        return None, "missing_blank"
    if isinstance(value, bool):
        return None, "invalid_boolean"
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 0 or int(value) != value:
            return None, "invalid_number"
        return int(value), "observed"
    if _text(value) == "-" and absence_marker_defined:
        return 0, "source_absence_marker"
    if _text(value) in {"…", "..."}:
        return None, "missing_source_marker"
    return None, "unparsed_text"


def _code(value):
    if value is None or _text(value) == "":
        return None, "missing"
    if isinstance(value, bool):
        return None, "invalid_boolean"
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 0 or int(value) != value:
            return None, "invalid_numeric_code"
        if value >= 10**15:
            return None, "numeric_precision_risk"
        return str(int(value)), "numeric_leading_zeros_unverified"
    # Internal whitespace is preserved: repairing it requires a classifier.
    value = str(value).strip()
    return (value, "text_digits") if re.fullmatch(r"[0-9]+", value) else (None, "unparsed_text")


def _row_kind(name: str, region_names: dict) -> str:
    if name in region_names:
        return "region_header"
    lower = name.lower()
    if "область" in lower and ("без " in lower or "кроме " in lower):
        return "region_scope_variant"
    if any(term in lower for term in ("муниципальн", "поселение", "городской округ")):
        return "municipality_explicit_label"
    if re.match(r"^(г|пгт|рп|гп|дп|кп|п)[ .]", lower):
        return "locality_explicit_label"
    return "territory_level_unresolved"


def parse_population_workbook(
    path: str | Path, *, source_url: str, retrieved_at: str
) -> tuple[pd.DataFrame, dict]:
    """Return all numeric territory rows plus a source audit, with no external joins.

    Date and persons units must be present in the workbook header. Blanks stay
    missing. A dash becomes zero only when the preface explicitly defines it as
    'явление отсутствует'. Original values and statuses remain in the output.
    Region context comes from section names listed within the same workbook;
    it is neither a verified external region code nor a boundary crosswalk.
    """
    path = Path(path)
    date.fromisoformat(retrieved_at)
    if not source_url.startswith("https://"):
        raise ValueError("A primary HTTPS source URL is required")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook[SHEET]
        header_cells = [cell for row in sheet.iter_rows(max_row=8) for cell in row if cell.value]
        header = " ".join(_text(cell.value) for cell in header_cells)
        matches = re.findall(r"на 1 января (20\d{2}) года", header)
        if len(set(matches)) != 1:
            raise ValueError("A unique reference date must be explicit in the table header")
        reference_date = f"{matches[0]}-01-01"
        if "(человек)" not in header:
            raise ValueError("Population units must be explicitly stated as persons")
        with_codes = "ТЕРСОН-МО" in header
        offset = 1 if with_codes else 0
        expected = ["Все", "население", "городское", "сельское"]
        if not all(term in header for term in expected):
            raise ValueError("Unrecognized population columns")
        preface = workbook["Предисловие"]
        preface_rows = [
            {"cell": cell.coordinate, "text": str(cell.value)}
            for row in preface.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        ]
        absence_defined = any(
            re.search(r"^-\s+-\s+явление отсутствует", item["text"].strip())
            for item in preface_rows
        )
        regions = {}
        for row in workbook[REGIONS_SHEET].iter_rows(min_row=2):
            if isinstance(row[0].value, (int, float)) and row[1].value:
                name = _text(row[1].value)
                if name in regions:
                    raise ValueError(f"Duplicate name in the source region directory: {name}")
                regions[name] = int(row[0].value)
        if not regions:
            raise ValueError("No source region directory entries")
        records, skipped = [], []
        region_name = region_row = region_ordinal = None
        start = 8 if with_codes else 9
        for row in sheet.iter_rows(min_row=start):
            name_cell = row[offset]
            name = _text(name_cell.value)
            cells = row[offset + 1 : offset + 4]
            if not name and all(cell.value is None for cell in cells):
                continue
            if name in regions:
                region_name, region_row, region_ordinal = name, name_cell.row, regions[name]
            is_section = name.rstrip().endswith(":") or name.lower().startswith("в том числе")
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(
                    f"Population formula requires explicit evaluation: row {name_cell.row}"
                )
            has_value = any(isinstance(cell.value, (int, float)) for cell in cells)
            has_marker = any(_text(cell.value) in {"-", "…", "..."} for cell in cells)
            if is_section or not (has_value or has_marker):
                skipped.append(
                    {
                        "source_row": name_cell.row,
                        "name_raw": str(name_cell.value),
                        "reason": "section_or_non_observation_text",
                    }
                )
                continue
            if not name:
                raise ValueError(f"Numeric row lacks territory name at {name_cell.row}")
            record = {
                "source_record_id": f"{digest}:{SHEET}:{name_cell.row}",
                "reference_date": reference_date,
                "unit": "persons",
                "region_name_source": region_name,
                "region_section_row": region_row,
                "source_region_ordinal": region_ordinal,
                "region_context_status": "source_section" if region_name else "unresolved",
                "territory_name": name,
                "source_name_raw": str(name_cell.value),
                "row_kind": _row_kind(name, regions),
                "source_indent": name_cell.alignment.indent,
                "source_name_bold": bool(name_cell.font.bold),
                "source_code_raw": str(row[0].value)
                if with_codes and row[0].value is not None
                else None,
                "source_code_system": "TERSON-MO" if with_codes else "not_provided",
                "oktmo": None,
                "territory_id": None,
                "join_status": "unlinked_no_verified_crosswalk",
                "source_filename": path.name,
                "source_sheet": SHEET,
                "source_row": name_cell.row,
                "source_value_range": f"{cells[0].coordinate}:{cells[-1].coordinate}",
                "source_url": source_url,
                "source_sha256": digest,
                "retrieved_at": retrieved_at,
                "published_at": None,
            }
            record["source_code"], record["source_code_status"] = (
                _code(row[0].value) if with_codes else (None, "not_provided")
            )
            for field, cell in zip(POPULATION_FIELDS, cells, strict=True):
                value, status = _population_value(
                    cell.value, absence_marker_defined=absence_defined
                )
                record[field], record[f"{field}_status"] = value, status
                record[f"{field}_raw"] = str(cell.value) if cell.value is not None else None
            values = [record[field] for field in POPULATION_FIELDS]
            if any(value is None for value in values):
                record["balance_status"] = "not_assessed_missing_component"
            else:
                record["balance_status"] = (
                    "pass" if values[0] == values[1] + values[2] else "mismatch"
                )
            record["urban_share"] = (
                record["population_urban"] / record["population_total"]
                if record["balance_status"] == "pass" and record["population_total"] > 0
                else None
            )
            records.append(record)
        frame = pd.DataFrame(records)
        if frame.empty:
            raise ValueError("No population observations found")
        for field in (*POPULATION_FIELDS, "oktmo", "territory_id"):
            frame[field] = frame[field].astype("string" if field == "oktmo" else "Int64")
        frame["source_code_duplicate"] = frame.source_code.notna() & frame.source_code.duplicated(
            keep=False
        )
        qa = {
            "status": "ready_unlinked_with_quality_flags",
            "source_filename": path.name,
            "source_url": source_url,
            "source_sha256": digest,
            "source_bytes": path.stat().st_size,
            "reference_date": reference_date,
            "retrieved_at": retrieved_at,
            "published_at": None,
            "header_evidence": [
                {"cell": cell.coordinate, "text": str(cell.value)} for cell in header_cells
            ],
            "preface_evidence": preface_rows,
            "sheet": SHEET,
            "n_rows": len(frame),
            "n_region_sections": int(frame.region_section_row.nunique()),
            "region_context_missing_rows": int(frame.region_name_source.isna().sum()),
            "row_kinds": frame.row_kind.value_counts().to_dict(),
            "source_code_system": "TERSON-MO" if with_codes else "not_provided",
            "source_code_status_counts": frame.source_code_status.value_counts().to_dict(),
            "duplicate_source_code_rows": int(frame.source_code_duplicate.sum()),
            "population_status_counts": {
                field: dict(Counter(frame[f"{field}_status"])) for field in POPULATION_FIELDS
            },
            "balance_counts": frame.balance_status.value_counts().to_dict(),
            "balance_mismatch_rows": frame.loc[
                frame.balance_status.eq("mismatch"), "source_row"
            ].tolist(),
            "source_absence_marker_is_zero": absence_defined,
            "all_rows_unlinked": True,
            "oktmo_provided": False,
            "skipped_rows": skipped,
            "warning": (
                "Mixed and overlapping territorial levels; do not sum all rows. "
                "Region context is not an external crosswalk."
            ),
        }
        return frame, qa
    finally:
        workbook.close()
